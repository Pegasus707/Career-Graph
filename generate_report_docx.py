import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def build_academic_report():
    doc = Document()

    # Page Margins (Standard 1 inch / 2.54 cm report margins)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)

    # Palette: Clean Academic / Corporate
    NAVY = RGBColor(30, 58, 138)       # #1E3A8A - Headings
    DARK_SLATE = RGBColor(30, 41, 59)  # #1E293B - Subheadings
    BODY_COLOR = RGBColor(30, 41, 59)  # #1E293B - Body text
    MUTED_COLOR = RGBColor(100, 116, 139)

    # Base Normal Style
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Calibri'
    style_normal.font.size = Pt(11)
    style_normal.font.color.rgb = BODY_COLOR
    style_normal.paragraph_format.line_spacing = 1.2
    style_normal.paragraph_format.space_after = Pt(5)

    # Helper: Set Cell XML Background
    def set_cell_background(cell, hex_color):
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
        cell._tc.get_or_add_tcPr().append(shd)

    # Helper: Set Cell Margins (Padding)
    def set_cell_padding(cell, top=120, bottom=120, left=160, right=160):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = parse_xml(
            f'<w:tcMar {nsdecls("w")}>'
            f'<w:top w:w="{top}" w:type="dxa"/>'
            f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
            f'<w:left w:w="{left}" w:type="dxa"/>'
            f'<w:right w:w="{right}" w:type="dxa"/>'
            f'</w:tcMar>'
        )
        tcPr.append(tcMar)

    # Helper: Clean Professional Table Styling
    def format_report_table(table, col_widths, has_header=True):
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        tblPr = table._tbl.tblPr
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="6" w:space="0" w:color="94A3B8"/>'
            f'<w:bottom w:val="single" w:sz="8" w:space="0" w:color="94A3B8"/>'
            f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
            f'<w:insideV w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
            f'<w:left w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
            f'<w:right w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
            f'</w:tblBorders>'
        )
        tblPr.append(borders)

        for row_idx, row in enumerate(table.rows):
            trPr = row._tr.get_or_add_trPr()
            trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

            if row_idx == 0 and has_header:
                trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

            for col_idx, cell in enumerate(row.cells):
                cell.width = col_widths[col_idx]
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                set_cell_padding(cell, top=100, bottom=100, left=140, right=140)

                if row_idx == 0 and has_header:
                    set_cell_background(cell, "1E3A8A")
                    for p in cell.paragraphs:
                        p.paragraph_format.space_before = Pt(2)
                        p.paragraph_format.space_after = Pt(2)
                        for r in p.runs:
                            r.font.bold = True
                            r.font.color.rgb = RGBColor(255, 255, 255)
                            r.font.size = Pt(9.5)
                else:
                    if row_idx % 2 == 1:
                        set_cell_background(cell, "FFFFFF")
                    else:
                        set_cell_background(cell, "F8FAFC")
                    for p in cell.paragraphs:
                        p.paragraph_format.space_before = Pt(1)
                        p.paragraph_format.space_after = Pt(1)
                        for r in p.runs:
                            r.font.size = Pt(9)
                            r.font.color.rgb = BODY_COLOR

    # Typography Helpers
    def add_chapter_title(num_str, title_str):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(20)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        
        r_num = p.add_run(f"CHAPTER {num_str}\n")
        r_num.bold = True
        r_num.font.size = Pt(12)
        r_num.font.color.rgb = MUTED_COLOR

        r_title = p.add_run(title_str)
        r_title.bold = True
        r_title.font.size = Pt(16)
        r_title.font.color.rgb = NAVY

        # Subtle bottom line under chapter title
        pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="8" w:space="2" w:color="1E3A8A"/></w:pBdr>')
        p._p.get_or_add_pPr().append(pBdr)

        p_spacer = doc.add_paragraph()
        p_spacer.paragraph_format.space_after = Pt(4)
        return p

    def add_section(num_str, text_str):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(f"{num_str}  {text_str}")
        r.bold = True
        r.font.size = Pt(12.5)
        r.font.color.rgb = DARK_SLATE
        return p

    def add_subsection(num_str, text_str):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(f"{num_str}  {text_str}")
        r.bold = True
        r.font.size = Pt(11)
        r.font.color.rgb = NAVY
        return p

    def add_bullet(bold_prefix, text_str):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            rb = p.add_run(bold_prefix)
            rb.bold = True
            rb.font.color.rgb = DARK_SLATE
        rt = p.add_run(text_str)
        rt.font.color.rgb = BODY_COLOR
        return p

    def add_note_box(label, text):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.cell(0, 0)
        cell.width = Inches(6.5)
        set_cell_background(cell, "F1F5F9")
        set_cell_padding(cell, top=120, bottom=120, left=160, right=160)
        
        tcPr = cell._tc.get_or_add_tcPr()
        tcBorders = parse_xml(
            f'<w:tcBorders {nsdecls("w")}>'
            f'<w:left w:val="single" w:sz="20" w:space="0" w:color="1E3A8A"/>'
            f'<w:top w:val="none"/>'
            f'<w:right w:val="none"/>'
            f'<w:bottom w:val="none"/>'
            f'</w:tcBorders>'
        )
        tcPr.append(tcBorders)

        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        rl = p.add_run(f"{label}: ")
        rl.bold = True
        rl.font.size = Pt(9.5)
        rl.font.color.rgb = NAVY
        rt = p.add_run(text)
        rt.font.size = Pt(9.5)
        rt.font.color.rgb = BODY_COLOR

        p_sp = doc.add_paragraph()
        p_sp.paragraph_format.space_after = Pt(4)

    # -------------------------------------------------------------
    # Setup Running Header & Footer
    # -------------------------------------------------------------
    header = doc.sections[0].header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hrun = hp.add_run("CareerGraph — Project Report")
    hrun.font.size = Pt(8.5)
    hrun.font.color.rgb = MUTED_COLOR

    footer = doc.sections[0].footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    frun1 = fp.add_run("Page ")
    frun1.font.size = Pt(8.5)
    frun1.font.color.rgb = MUTED_COLOR
    fldSimple = parse_xml(r'<w:fldSimple %s w:instr="PAGE"/>' % nsdecls('w'))
    fp._p.append(fldSimple)

    # =============================================================
    # COVER PAGE (Formal Project Report Format)
    # =============================================================
    cp1 = doc.add_paragraph()
    cp1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp1.paragraph_format.space_before = Pt(24)
    cp1.paragraph_format.space_after = Pt(8)
    r = cp1.add_run("A PROJECT REPORT ON")
    r.bold = True
    r.font.size = Pt(13)
    r.font.color.rgb = MUTED_COLOR

    cp_title = doc.add_paragraph()
    cp_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp_title.paragraph_format.space_before = Pt(12)
    cp_title.paragraph_format.space_after = Pt(12)
    rt = cp_title.add_run("CAREERGRAPH: AN INTERACTIVE AND PERSONALIZED IT CAREER ROADMAP PLATFORM WITH PREREQUISITE LOCKING AND SKILL VERIFICATION")
    rt.bold = True
    rt.font.size = Pt(18)
    rt.font.color.rgb = NAVY

    cp_sub = doc.add_paragraph()
    cp_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp_sub.paragraph_format.space_before = Pt(8)
    cp_sub.paragraph_format.space_after = Pt(28)
    rsub = cp_sub.add_run("Submitted in partial fulfillment of the requirements for the Degree of\nBachelor of Engineering / Technology in Computer Science & Information Technology")
    rsub.font.size = Pt(11)
    rsub.font.italic = True
    rsub.font.color.rgb = DARK_SLATE

    # Authors & Details Box
    t_cover = doc.add_table(rows=5, cols=2)
    t_cover.alignment = WD_TABLE_ALIGNMENT.CENTER
    cov_widths = [Inches(2.5), Inches(4.0)]
    cov_data = [
        ("Project Title:", "CareerGraph"),
        ("Project Candidates:", "Archit Surve\nSanika Vichare"),
        ("Project Guide / Mentor:", "Project Guide / Faculty Advisor"),
        ("Academic Year:", "2025 – 2026"),
        ("Department & College:", "Department of Computer Engineering / IT")
    ]
    for idx, (lbl, val) in enumerate(cov_data):
        row = t_cover.rows[idx]
        cell_l, cell_r = row.cells[0], row.cells[1]
        cell_l.width, cell_r.width = cov_widths[0], cov_widths[1]
        set_cell_padding(cell_l, 60, 60, 100, 100)
        set_cell_padding(cell_r, 60, 60, 100, 100)
        set_cell_background(cell_l, "F8FAFC")
        set_cell_background(cell_r, "FFFFFF")
        
        pl = cell_l.paragraphs[0]
        rl = pl.add_run(lbl)
        rl.bold = True
        rl.font.size = Pt(10)
        rl.font.color.rgb = NAVY

        pr = cell_r.paragraphs[0]
        rr = pr.add_run(val)
        rr.font.size = Pt(10)
        rr.font.color.rgb = DARK_SLATE

    format_report_table(t_cover, cov_widths, has_header=False)

    doc.add_page_break()

    # =============================================================
    # CERTIFICATE & ACKNOWLEDGEMENT
    # =============================================================
    cert_p = doc.add_paragraph()
    cert_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cert_p.paragraph_format.space_before = Pt(16)
    cert_p.paragraph_format.space_after = Pt(12)
    rcert = cert_p.add_run("CERTIFICATE")
    rcert.bold = True
    rcert.font.size = Pt(16)
    rcert.font.color.rgb = NAVY

    doc.add_paragraph(
        "This is to certify that the project report entitled \"CareerGraph: An Interactive and Personalized IT "
        "Career Roadmap Platform with Prerequisite Locking and Skill Verification\" is a bona fide record of work "
        "carried out by Archit Surve and Sanika Vichare in partial fulfillment of the requirements for the "
        "award of the Bachelor's Degree in Computer Engineering / Information Technology."
    )
    doc.add_paragraph(
        "The work embodied in this project report has been completed under our supervision and guidance, "
        "and complies with the prescribed academic guidelines and standards."
    )

    p_sig = doc.add_paragraph()
    p_sig.paragraph_format.space_before = Pt(40)
    p_sig.paragraph_format.space_after = Pt(20)
    rsig = p_sig.add_run("_________________________\t\t\t_________________________\nProject Guide / Mentor\t\t\tHead of Department (HOD)")
    rsig.bold = True
    rsig.font.size = Pt(10)
    rsig.font.color.rgb = DARK_SLATE

    # Acknowledgement
    ack_p = doc.add_paragraph()
    ack_p.paragraph_format.space_before = Pt(24)
    ack_p.paragraph_format.space_after = Pt(8)
    rack = ack_p.add_run("ACKNOWLEDGEMENT")
    rack.bold = True
    rack.font.size = Pt(13)
    rack.font.color.rgb = NAVY

    doc.add_paragraph(
        "We express our sincere gratitude to our Project Guide and faculty members for their constant encouragement, "
        "insightful advice, and constructive feedback throughout the design, development, and deployment of CareerGraph. "
        "We are also thankful to our institution for providing the infrastructure, computing resources, and academic "
        "environment necessary to carry out this full-stack engineering project."
    )

    doc.add_page_break()

    # =============================================================
    # ABSTRACT
    # =============================================================
    abs_p = doc.add_paragraph()
    abs_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    abs_p.paragraph_format.space_before = Pt(16)
    abs_p.paragraph_format.space_after = Pt(12)
    rabs = abs_p.add_run("ABSTRACT")
    rabs.bold = True
    rabs.font.size = Pt(16)
    rabs.font.color.rgb = NAVY

    doc.add_paragraph(
        "In the current technical education and software industry landscape, students and career-switchers encounter "
        "an overwhelming volume of tutorials, courses, and documentation. However, learners frequently struggle because "
        "they lack structured sequencing: foundational prerequisites (such as Git and Linux) are often skipped before "
        "attempting advanced frameworks and cloud architectures. Furthermore, existing roadmap platforms (such as roadmap.sh) "
        "are largely static and cannot assess what a learner already knows, nor do they carry earned progress across adjacent career paths."
    )
    doc.add_paragraph(
        "To solve these challenges, we designed and implemented CareerGraph, an interactive full-stack web application "
        "built with Vue 3, Node.js, Express, and MongoDB Atlas. CareerGraph introduces four core innovations: "
        "(1) a 5-tier learning hierarchy (Career Track → Skill → Course → Level/Module → Lesson), "
        "(2) a topological prerequisite locking engine that organizes learning into Foundations, Core Stack, and Advanced phases, "
        "(3) universal cross-track skill progress reflection, ensuring that foundational skills (e.g., Git or Python) mastered in one track automatically credit towards other tracks, and "
        "(4) a server-side Proof-of-Skill validation quiz engine that unlocks a verified badge upon passing."
    )
    doc.add_paragraph(
        "The system has been fully deployed with an interactive SVG Bézier canvas on Vercel Edge and a RESTful API on Render. "
        "The application provides learners with a structured, step-by-step path to IT career readiness while offering educators "
        "a transparent mechanism to track student skill acquisition."
    )

    doc.add_page_break()

    # =============================================================
    # TABLE OF CONTENTS
    # =============================================================
    toc_p = doc.add_paragraph()
    toc_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    toc_p.paragraph_format.space_before = Pt(16)
    toc_p.paragraph_format.space_after = Pt(14)
    rtoc = toc_p.add_run("TABLE OF CONTENTS")
    rtoc.bold = True
    rtoc.font.size = Pt(16)
    rtoc.font.color.rgb = NAVY

    toc_items = [
        ("Certificate & Acknowledgement", "ii"),
        ("Abstract", "iii"),
        ("Chapter 1: Introduction", "1"),
        ("    1.1 Background & Motivation", "1"),
        ("    1.2 Problem Statement", "1"),
        ("    1.3 Project Objectives", "2"),
        ("    1.4 Scope of the Project", "2"),
        ("Chapter 2: Literature Survey & Existing System Analysis", "3"),
        ("    2.1 Review of Existing Platforms", "3"),
        ("    2.2 Drawbacks of Current Systems", "3"),
        ("    2.3 Proposed System Highlights", "4"),
        ("    2.4 Feature Comparison Table", "4"),
        ("Chapter 3: System Requirements & Architecture", "5"),
        ("    3.1 Hardware and Software Requirements", "5"),
        ("    3.2 Three-Tier Decoupled Architecture", "5"),
        ("    3.3 Database Modeling & Entity Relationship Diagram", "7"),
        ("    3.4 Database Collections Summary", "8"),
        ("Chapter 4: Implementation of Core Modules", "9"),
        ("    4.1 Module 1: User Authentication & Onboarding Assessment", "9"),
        ("    4.2 Module 2: Topological Prerequisite & Phase Locking Engine", "9"),
        ("    4.3 Module 3: Learning Curriculum & Progress Aggregation", "10"),
        ("    4.4 Module 4: Proof-of-Skill Quiz & Server-Side Verification", "11"),
        ("    4.5 Module 5: Universal Cross-Track Skill Sharing", "12"),
        ("Chapter 5: User Interface & Workflow Walkthrough", "13"),
        ("    5.1 User Registration and Career Goal Selection", "13"),
        ("    5.2 Interactive SVG Bézier Roadmap Canvas", "13"),
        ("    5.3 Topic Preview Drawer & Lesson Checklists", "14"),
        ("    5.4 Taking the Verification Quiz and Earning the Badge", "14"),
        ("    5.5 Switching Career Tracks On-The-Fly", "14"),
        ("Chapter 6: System Testing & Experimental Results", "15"),
        ("    6.1 Testing Methodology", "15"),
        ("    6.2 Functional Test Cases & Validation Results", "15"),
        ("Chapter 7: Conclusion & Future Scope", "17"),
        ("    7.1 Conclusion", "17"),
        ("    7.2 Future Enhancements", "17"),
        ("References", "18")
    ]

    t_toc = doc.add_table(rows=len(toc_items), cols=2)
    format_report_table(t_toc, [Inches(5.3), Inches(1.2)], has_header=False)
    for idx, (item, page) in enumerate(toc_items):
        row = t_toc.rows[idx]
        cell_item, cell_page = row.cells[0], row.cells[1]
        set_cell_padding(cell_item, 30, 30, 60, 60)
        set_cell_padding(cell_page, 30, 30, 60, 60)
        set_cell_background(cell_item, "FFFFFF")
        set_cell_background(cell_page, "FFFFFF")

        p_it = cell_item.paragraphs[0]
        r_it = p_it.add_run(item)
        r_it.font.size = Pt(9.5)
        if "Chapter" in item or "Abstract" in item or "Certificate" in item or "References" in item:
            r_it.bold = True
            r_it.font.color.rgb = NAVY
        else:
            r_it.font.color.rgb = DARK_SLATE

        p_pg = cell_page.paragraphs[0]
        p_pg.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r_pg = p_pg.add_run(page)
        r_pg.font.size = Pt(9.5)
        r_pg.font.color.rgb = MUTED_COLOR

    doc.add_page_break()

    # =============================================================
    # CHAPTER 1: INTRODUCTION
    # =============================================================
    add_chapter_title("1", "INTRODUCTION")

    add_section("1.1", "Background & Motivation")
    doc.add_paragraph(
        "Over the last decade, technical education has witnessed a dramatic shift from traditional textbooks "
        "to open online tutorials, video platforms, and bootcamps. While learning materials are widely available, "
        "engineering students and aspiring software professionals face a major challenge known as the 'paradox of choice'. "
        "Learners struggle to determine what to study, in what order, and how their current knowledge aligns with real industry expectations."
    )
    doc.add_paragraph(
        "Many students attempt to learn high-level frameworks like React or Docker before mastering fundamentals like "
        "HTML/CSS layout concepts, JavaScript event loops, or command-line Git operations. Skipping core fundamentals "
        "leads to conceptual confusion, fragile projects, and high abandonment rates. CareerGraph was developed to "
        "provide a structured, interactive guidance system that enforces foundational prerequisites before unlocking advanced topics."
    )

    add_section("1.2", "Problem Statement")
    doc.add_paragraph("Specifically, modern learners face four critical obstacles:")
    add_bullet("Absence of Prerequisite Enforcement: ", "Online courses present modules in isolation. A student can enroll directly in a cloud deployment course without ever touching Linux commands or networking fundamentals.")
    add_bullet("Static and Read-Only Roadmaps: ", "Existing community roadmaps provide static PNG/SVG diagrams. They show an ideal curriculum but cannot adapt to what an individual student already knows.")
    add_bullet("Siloed, Disjointed Progress: ", "When a student studies a skill like Python or SQL in one learning path, their effort is lost if they decide to switch tracks (e.g., from Full Stack Web Development to Data Analysis).")
    add_bullet("Lack of Credible Verification: ", "Most e-learning platforms rely on unchecked self-reporting (ticking checkboxes) without evaluating whether the student truly understands the underlying principles.")

    add_section("1.3", "Project Objectives")
    doc.add_paragraph("The primary objectives of CareerGraph are:")
    add_bullet("1. Structured Hierarchical Curriculum: ", "Organize IT education into a 5-tier structure: Career Track → Skill → Course → Level/Module → Lesson.")
    add_bullet("2. Topological Prerequisite Engine: ", "Automatically lock advanced learning phases until preceding foundation skills are marked as completed.")
    add_bullet("3. Universal Cross-Track Credit: ", "Recognize skills via normalized global identifiers so that any skill mastered in one career track instantly reflects across all other tracks.")
    add_bullet("4. Proof-of-Skill Validation: ", "Provide server-graded micro-quizzes for skills to verify authentic understanding before awarding verified badges.")
    add_bullet("5. Interactive Visual Roadmap Canvas: ", "Deliver an intuitive web interface with dynamic Bézier connector curves, pan/zoom controls, topic drawers, and instant status updates.")

    add_section("1.4", "Scope of the Project")
    doc.add_paragraph(
        "CareerGraph focuses on major Computer Science and Information Technology career tracks: "
        "Full Stack Developer, DevOps Engineer, AI/ML Engineer, Frontend Developer, Backend Developer, and Data Analyst. "
        "The project includes complete onboarding questionnaires, dynamic graph generation, interactive checklists, "
        "server-side quiz evaluation, and live cloud deployment on Vercel and Render."
    )

    doc.add_page_break()

    # =============================================================
    # CHAPTER 2: LITERATURE SURVEY
    # =============================================================
    add_chapter_title("2", "LITERATURE SURVEY & EXISTING SYSTEM ANALYSIS")

    add_section("2.1", "Review of Existing Platforms")
    doc.add_paragraph(
        "To establish the context of CareerGraph, we analyzed existing tools used by students and industry developers:"
    )
    add_bullet("roadmap.sh: ", "A popular community-driven platform providing comprehensive visual roadmaps for various developer roles. While it serves as an excellent curriculum reference, it functions primarily as a static guide. It does not enforce prerequisites, offer integrated interactive courses, or provide automated knowledge verification.")
    add_bullet("Coursera & Udemy: ", "Massive Open Online Course (MOOC) platforms offering structured video content. However, courses are standalone purchases. They lack an overarching career graph connecting disparate tools, and cross-course skill tracking is nonexistent.")
    add_bullet("LeetCode & HackerRank: ", "Platforms focused purely on algorithmic problem-solving and coding interviews. They do not guide a student through holistic software engineering domains like system design, DevOps tooling, or web architecture.")

    add_section("2.2", "Drawbacks of Current Systems")
    doc.add_paragraph(
        "Existing systems either present roadmaps without active tracking, or provide courses without a career-level "
        "dependency graph. None of them dynamically evaluate a student's prior degree/stream background, compute "
        "their remaining skill gap, lock dependent phases, or universally credit shared competencies."
    )

    add_section("2.3", "Proposed CareerGraph Highlights")
    doc.add_paragraph(
        "CareerGraph bridges the gap between static reference roadmaps and interactive learning portals. "
        "It introduces dynamic roadmap assembly, visual graph navigation, and server-side proof-of-skill validation."
    )

    add_section("2.4", "Feature Comparison Table")
    comp_data = [
        ("Feature / Capability", "roadmap.sh", "Standard MOOCs", "CareerGraph (Proposed)"),
        ("Visual Career Roadmap", "Static SVG / Community", "Linear Syllabus List", "Interactive SVG Bézier Canvas"),
        ("Prerequisite Locking", "None (All visible)", "Course-specific only", "Topological Phase & Node Locking"),
        ("Personalized Gap Analysis", "None", "Basic quiz placement", "Onboarding Evaluation + Dynamic Calc"),
        ("Cross-Track Skill Credit", "Not Supported", "Not Supported", "Universal skillId Propagation"),
        ("Verified Proof-of-Skill", "Honor system checkmarks", "Certificate on video completion", "Server-Graded Validation Quiz"),
        ("In-Place Career Switching", "Requires page navigation", "Separate course enrollment", "Instant modal switch with live update")
    ]
    t_comp = doc.add_table(rows=len(comp_data), cols=4)
    format_report_table(t_comp, [Inches(1.8), Inches(1.5), Inches(1.6), Inches(1.6)])
    for r_idx, row in enumerate(t_comp.rows):
        for c_idx, cell in enumerate(row.cells):
            cell.text = comp_data[r_idx][c_idx]

    doc.add_page_break()

    # =============================================================
    # CHAPTER 3: SYSTEM REQUIREMENTS & ARCHITECTURE
    # =============================================================
    add_chapter_title("3", "SYSTEM REQUIREMENTS & ARCHITECTURE")

    add_section("3.1", "Hardware and Software Requirements")
    doc.add_paragraph("The system was designed, developed, and tested under the following environment specifications:")

    req_data = [
        ("Component", "Development / Server Requirement", "Client / End-User Requirement"),
        ("Operating System", "macOS Sonoma / Linux Ubuntu 22.04 / Windows 11", "Any OS with a modern web browser"),
        ("Processor", "Multi-core 64-bit CPU (Intel i5/M1 or higher)", "Dual-core processor or modern smartphone"),
        ("Memory (RAM)", "8 GB RAM minimum (16 GB recommended)", "4 GB RAM minimum"),
        ("Runtime & Framework", "Node.js 18+, Express 4.19, Vite 5.3", "Chrome 100+, Firefox 100+, Safari 15+, Edge"),
        ("Database", "MongoDB Atlas (M0/M10) or local MongoDB 7.0+", "None (Accessed through REST API)")
    ]
    t_req = doc.add_table(rows=len(req_data), cols=3)
    format_report_table(t_req, [Inches(1.8), Inches(2.5), Inches(2.2)])
    for r_idx, row in enumerate(t_req.rows):
        for c_idx, cell in enumerate(row.cells):
            cell.text = req_data[r_idx][c_idx]

    add_section("3.2", "Three-Tier Decoupled Architecture")
    doc.add_paragraph(
        "CareerGraph adopts a clean three-tier architecture to maintain separation of concerns, "
        "facilitate team collaboration, and allow independent deployment and scaling."
    )

    if os.path.exists("docs/system_architecture.png"):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(8)
        p_img.paragraph_format.space_after = Pt(2)
        doc.add_picture("docs/system_architecture.png", width=Inches(5.8))
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_after = Pt(10)
        rcap = p_cap.add_run("Figure 3.1: Three-Tier System Architecture of CareerGraph")
        rcap.font.size = Pt(8.5)
        rcap.font.italic = True
        rcap.font.color.rgb = MUTED_COLOR

    doc.add_paragraph(
        "1. Presentation Tier (Frontend): Built with Vue 3 (Composition API) and Vite. Centralized state is managed "
        "using Pinia stores. It renders the interactive roadmap canvas using vector SVG Bézier connectors, hosts the "
        "topic preview drawer, and displays the quiz modal."
    )
    doc.add_paragraph(
        "2. Application Tier (Backend API): Built with Node.js and Express.js. Handles stateless JWT authentication, "
        "bcrypt password hashing, topological dependency resolution (skillGapService.js), and server-side quiz grading."
    )
    doc.add_paragraph(
        "3. Persistence Tier (Database): MongoDB managed via Mongoose 8.5 ODM. Normalized schemas represent entities, "
        "while compound uniqueness indexes ensure high concurrency integrity."
    )

    add_section("3.3", "Database Modeling & Entity Relationship Diagram")
    doc.add_paragraph(
        "The database design consists of 9 distinct collections. Content (Skills, Courses, Levels) is normalized "
        "to avoid data redundancy, while user progress records reference universal keys for cross-track sharing."
    )

    if os.path.exists("docs/er_diagram.png"):
        p_img2 = doc.add_paragraph()
        p_img2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img2.paragraph_format.space_before = Pt(8)
        p_img2.paragraph_format.space_after = Pt(2)
        doc.add_picture("docs/er_diagram.png", width=Inches(5.8))
        
        p_cap2 = doc.add_paragraph()
        p_cap2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap2.paragraph_format.space_after = Pt(10)
        rcap2 = p_cap2.add_run("Figure 3.2: Entity-Relationship Diagram and Schema Associations")
        rcap2.font.size = Pt(8.5)
        rcap2.font.italic = True
        rcap2.font.color.rgb = MUTED_COLOR

    add_section("3.4", "Database Collections Summary")
    models_summary = [
        ("Collection / Model", "Primary Fields & Key Types", "Role in System"),
        ("User", "_id, name, email (UK), password (hash), targetCareer (FK)", "Stores user identity and active career reference."),
        ("UserProfile", "_id, user (1:1 FK), education, status, skills array", "Holds academic stream, degree, and declared baseline ratings."),
        ("Career", "_id, name, slug (UK), streams, degrees, requiredSkills", "Defines career tracks and skills mapped to phases."),
        ("Skill", "_id, skillId (UK), name, prerequisites (self FK), resources", "Universal atomic skill catalog with dependency links."),
        ("Course", "_id, title, skill (1:1 FK), description", "Curriculum container connected to a specific skill."),
        ("Level", "_id, course (FK), name (Beg, Int, Adv), modules array", "Curriculum levels containing modules and atomic lessons."),
        ("CourseProgress", "_id, user (FK), course (FK), percent, verified", "Aggregates overall course percentage and verification flag."),
        ("LessonProgress", "_id, user (FK), lesson (FK), completedAt", "Atomic record of completed lessons with compound uniqueness."),
        ("SkillQuiz", "_id, skill (1:1 FK), questions array (min 3)", "Stores validation questions, options, and correct answers.")
    ]
    t_mod = doc.add_table(rows=len(models_summary), cols=3)
    format_report_table(t_mod, [Inches(1.8), Inches(2.5), Inches(2.2)])
    for r_idx, row in enumerate(t_mod.rows):
        for c_idx, cell in enumerate(row.cells):
            cell.text = models_summary[r_idx][c_idx]

    doc.add_page_break()

    # =============================================================
    # CHAPTER 4: IMPLEMENTATION OF CORE MODULES
    # =============================================================
    add_chapter_title("4", "IMPLEMENTATION OF CORE MODULES")

    add_section("4.1", "Module 1: User Authentication & Onboarding Assessment")
    doc.add_paragraph(
        "The onboarding process captures the student's academic background and current knowledge level before constructing "
        "their custom roadmap:"
    )
    add_bullet("Registration & JWT Security: ", "Passwords are encrypted using bcryptjs with 10 salt rounds in Mongoose pre('save') hooks. On login, the server issues a signed, stateless JSON Web Token (JWT) with a 7-day expiration.")
    add_bullet("Academic Stream & Degree Filter: ", "During onboarding, the student selects their educational degree (e.g., B.Tech, BCA, B.Sc) and stream (Computer Science, Information Technology, AI/DS). The system matches eligible career tracks.")
    add_bullet("Self-Assessment Baseline: ", "The student rates their initial familiarity with core skills on a 0 to 4 scale ('Never heard of it' to 'Know everything'). This establishes the baseline for the skill gap calculation.")

    add_section("4.2", "Module 2: Topological Prerequisite & Phase Locking Engine")
    doc.add_paragraph(
        "Implemented in backend/src/services/skillGapService.js, this module enforces logical curriculum sequencing:"
    )
    add_bullet("Phase 1: Foundations: ", "Covers essential prerequisite skills (e.g., HTML, CSS, JavaScript, Git for Full Stack; Linux, Git, Networking for DevOps). Phase 1 is always unlocked.")
    add_bullet("Phase 2: Core Stack: ", "Covers main development tools (e.g., Vue, Node.js, Express, MongoDB). Phase 2 remains strictly locked until every skill in Phase 1 has reached the completed status.")
    add_bullet("Phase 3: Advanced & Ecosystem: ", "Covers architecture and production tooling (e.g., Docker, AWS Cloud, Kubernetes). Phase 3 remains locked until all Phase 2 skills are completed.")
    add_bullet("Node-Level Dependency Resolution: ", "Even inside an unlocked phase, individual nodes verify their direct prerequisite array. For example, JavaScript requires HTML and CSS; if either prerequisite is incomplete, JavaScript is marked locked with a tooltip explaining which skills must be finished first.")
    add_bullet("Next Recommended Skill Algorithm: ", "The engine traverses unlocked, uncompleted skills in topological order and automatically assigns the 'recommended' attribute to the first eligible skill, guiding the learner directly to their next action.")

    add_section("4.3", "Module 3: Learning Curriculum & Progress Aggregation")
    doc.add_paragraph(
        "Each skill links to a comprehensive syllabus structured into Beginner, Intermediate, and Advanced levels:"
    )
    add_bullet("Atomic Lesson Checklists: ", "Students toggle individual lessons in the slide-out preview drawer. When clicked, the backend creates or removes a LessonProgress record.")
    add_bullet("Hierarchical Progress Rollup: ", "Whenever a lesson is toggled, recalculateCourseProgress calculates the completion ratio: percent = (completedLessons / totalLessons) * 100, and atomically updates CourseProgress.")
    add_bullet("Overall Career Progress: ", "The roadmap aggregates progress across all required skills in the career track to compute a global percentage displayed in the dashboard and canvas header.")

    add_section("4.4", "Module 4: Proof-of-Skill Quiz & Server-Side Verification")
    doc.add_paragraph(
        "To provide credible proof that a student has mastered a subject, CareerGraph incorporates an integrated "
        "validation quiz engine (quizController.js):"
    )
    add_bullet("Advanced Phase Unlocking Rule: ", "The quiz cannot be taken on day one. A student must reach the Advanced phase (by completing all Beginner and Intermediate lessons, or having a declared proficiency >= 3) before the quiz unlocks.")
    add_bullet("Cheating Prevention via Payload Sanitization: ", "When questions are fetched by the browser, the server strips correctIndex and explanation from the JSON payload. Students cannot inspect the answers via browser DevTools.")
    add_bullet("Server-Side Grading: ", "Answers are submitted to POST /api/skills/:skillId/quiz/submit. The server compares selections against database answer keys. A passing grade requires at least 2 out of 3 correct answers (>= 66.7%).")
    add_bullet("Verified Badge & Level Elevation: ", "Upon passing, the system awards the Verified Badge, sets CourseProgress to 100%, marks all lessons complete, and elevates the user's skill profile to Level 4 (Expert).")

    add_section("4.5", "Module 5: Universal Cross-Track Skill Sharing")
    doc.add_paragraph(
        "Because skills are tagged with universal strings (e.g., 'git', 'docker', 'python'), progress earned in one "
        "track carries over universally:"
    )
    add_bullet("Cross-Track Recognition: ", "If a student masters Git (100%) while following the Full Stack Developer track, and later switches to the DevOps Engineer or AI/ML Engineer track, Git is already marked 100% and verified on the new roadmap.")
    add_bullet("Eliminating Redundant Effort: ", "Students never re-do foundational lessons they have already completed, mirroring real-world software engineering practice.")

    doc.add_page_break()

    # =============================================================
    # CHAPTER 5: USER INTERFACE & WORKFLOW WALKTHROUGH
    # =============================================================
    add_chapter_title("5", "USER INTERFACE & WORKFLOW WALKTHROUGH")

    add_section("5.1", "User Registration and Career Goal Selection")
    doc.add_paragraph(
        "The student begins at the Signup screen, entering their name, email, and password. Upon registration, "
        "they are directed into the multi-step Onboarding wizard. Here they select their college degree, "
        "their academic stream, and their primary career goal. Baseline skill ratings are collected through an "
        "intuitive slider interface."
    )

    add_section("5.2", "Interactive SVG Bézier Roadmap Canvas")
    doc.add_paragraph(
        "Upon completing onboarding, the student enters the interactive Roadmap view (Roadmap.vue):"
    )
    add_bullet("Canvas Graph View: ", "Displays skills grouped into Phase 1 (Foundations), Phase 2 (Core Stack), and Phase 3 (Advanced). Dynamic SVG cubic Bézier curves connect prerequisite nodes to dependent nodes with directional arrowheads.")
    add_bullet("Pan and Zoom Navigation: ", "The learner can drag to pan across the canvas and use the mouse wheel or floating HUD buttons to zoom between 50% and 150%. A 'Center View' button resets the viewport instantly.")
    add_bullet("1-Click Status Toggling: ", "Clicking on any node's status badge cycles its status (Not Started ➔ In Progress ➔ Completed). A custom right-click context menu provides direct access to change statuses.")
    add_bullet("Visual Lock Indication: ", "Locked skills display a padlock icon (🔒) with a tooltip explaining unmet prerequisites. Clicking a locked skill triggers a clear banner toast.")

    add_section("5.3", "Topic Preview Drawer & Lesson Checklists")
    doc.add_paragraph(
        "Clicking any node slides open the SkillPreviewDrawer without leaving the canvas. "
        "The drawer displays: (1) what the skill is, (2) why it matters in industry, (3) curated free documentation "
        "and video tutorials, and (4) an interactive lesson checklist organized by Beginner, Intermediate, and Advanced levels."
    )

    add_section("5.4", "Taking the Verification Quiz and Earning the Badge")
    doc.add_paragraph(
        "When the student finishes the Beginner and Intermediate lessons, a 'Take Proof-of-Skill Quiz' button unlocks. "
        "Clicking it opens the SkillQuizModal. The student answers 3 targeted multiple-choice questions. "
        "Upon submission, the server grades the quiz, provides detailed question explanations, and unlocks the Verified Badge (🛡️)."
    )

    add_section("5.5", "Switching Career Tracks On-The-Fly")
    doc.add_paragraph(
        "Students are not locked into a single career. Clicking 'Switch / Choose Another Career' opens the CareerSwitchModal. "
        "The student can search through all tracks, filter by academic stream, and switch immediately. "
        "The roadmap recalculates dynamically, preserving all previously completed skills."
    )

    doc.add_page_break()

    # =============================================================
    # CHAPTER 6: SYSTEM TESTING & EXPERIMENTAL RESULTS
    # =============================================================
    add_chapter_title("6", "SYSTEM TESTING & EXPERIMENTAL RESULTS")

    add_section("6.1", "Testing Methodology")
    doc.add_paragraph(
        "A comprehensive testing strategy was executed covering unit verification, integration testing, "
        "security testing, and user acceptance testing across both the Node.js backend and Vue 3 frontend."
    )

    add_section("6.2", "Functional Test Cases & Validation Results")
    test_cases = [
        ("TC ID", "Test Scenario", "Input / Action", "Expected Output", "Status"),
        ("TC-01", "User Registration & Security", "Enter valid name, email, password", "User created; password hashed with bcrypt (10 rounds); JWT returned", "PASS"),
        ("TC-02", "Duplicate Registration", "Register with existing email", "HTTP 400 error: 'User already exists'", "PASS"),
        ("TC-03", "Onboarding Skill Gap Calc", "Select Full Stack, rate HTML=3, JS=1", "HTML marked 'completed'; JS marked 'in_progress'; others 'not_started'", "PASS"),
        ("TC-04", "Phase 2 Gating Enforcement", "Attempt to toggle Phase 2 node with Phase 1 incomplete", "Node remains locked; banner displays: 'Complete all Phase 1 skills first'", "PASS"),
        ("TC-05", "Direct Node Prerequisite Lock", "Attempt to complete CSS before HTML", "CSS locked; tooltip: 'Prerequisite HTML must be completed first'", "PASS"),
        ("TC-06", "Canvas 1-Click Status Toggle", "Click node status badge on canvas", "Status cycles immediately; backend API updates; overall progress updates", "PASS"),
        ("TC-07", "Quiz Access Authorization", "Attempt GET /api/skills/:id/quiz before reaching Advanced", "HTTP 400 error: 'Must reach Advanced phase before taking verification quiz'", "PASS"),
        ("TC-08", "Quiz Payload Sanitization", "Inspect GET /api/skills/:id/quiz JSON", "Questions and options present; correctIndex and explanation stripped", "PASS"),
        ("TC-09", "Server-Side Quiz Grading", "Submit 2 out of 3 correct answers", "Score=2/3; passed=true; Verified Badge awarded; course progress set to 100%", "PASS"),
        ("TC-10", "Quiz Failure Handling", "Submit 1 out of 3 correct answers", "Score=1/3; passed=false; badge withheld; review explanations shown", "PASS"),
        ("TC-11", "Cross-Track Skill Credit", "Switch from Full Stack to DevOps with Git completed", "Git displays 100% and Verified in DevOps roadmap immediately", "PASS"),
        ("TC-12", "Production SPA Routing", "Refresh browser on /roadmap on Vercel", "vercel.json rewrite rules load index.html without 404 error", "PASS")
    ]
    t_test = doc.add_table(rows=len(test_cases), cols=5)
    format_report_table(t_test, [Inches(0.7), Inches(1.6), Inches(1.8), Inches(1.8), Inches(0.6)])
    for r_idx, row in enumerate(t_test.rows):
        for c_idx, cell in enumerate(row.cells):
            cell.text = test_cases[r_idx][c_idx]

    doc.add_page_break()

    # =============================================================
    # CHAPTER 7: CONCLUSION & FUTURE SCOPE
    # =============================================================
    add_chapter_title("7", "CONCLUSION & FUTURE SCOPE")

    add_section("7.1", "Conclusion")
    doc.add_paragraph(
        "CareerGraph successfully addresses the challenges of unstructured and fragmented IT education. "
        "By implementing a 5-tier learning hierarchy and a topological prerequisite locking engine, "
        "the platform ensures that students master foundational skills before venturing into advanced production technologies. "
        "The inclusion of universal cross-track skill sharing prevents redundant work, while server-graded Proof-of-Skill "
        "micro-quizzes provide genuine validation of student capabilities."
    )
    doc.add_paragraph(
        "The project has been engineered with modern full-stack best practices: Vue 3 Composition API with Vite, "
        "Pinia centralized state management, responsive SVG Bézier canvas navigation, Node.js/Express REST APIs with "
        "stateless JWT security, and a MongoDB Atlas persistence layer. Live deployments on Vercel and Render confirm "
        "that the platform is robust, performant, and production-ready."
    )

    add_section("7.2", "Future Enhancements")
    doc.add_paragraph("Potential future directions to expand CareerGraph include:")
    add_bullet("1. AI-Driven Adaptive Roadmaps: ", "Integrating Large Language Models (LLMs) to automatically generate personalized lesson explanations and real-time coding hints based on student quiz mistakes.")
    add_bullet("2. Community Roadmap Creator: ", "Allowing verified industry experts and senior mentors to publish custom, peer-reviewed roadmaps for emerging technologies (e.g., Web3, Rust, Edge AI).")
    add_bullet("3. Interactive In-Browser Code Sandboxes: ", "Embedding WebAssembly (Wasm) coding playgrounds directly into the topic drawer so learners can run code examples without leaving the platform.")
    add_bullet("4. Employer Talent Portal: ", "Enabling verified companies to search for candidates who have earned verified badges in specific prerequisite stacks.")

    # References
    add_section("7.3", "References")
    references = [
        "1. roadmap.sh — Community-driven developer roadmaps and guides (https://roadmap.sh).",
        "2. Vue.js 3 Official Documentation — The Progressive JavaScript Framework (https://vuejs.org).",
        "3. Node.js Foundation — Event-driven asynchronous JavaScript runtime (https://nodejs.org).",
        "4. Express.js Documentation — Fast, unopinionated, minimalist web framework (https://expressjs.com).",
        "5. MongoDB Manual — Document-based distributed database and ODM specifications (https://www.mongodb.com/docs).",
        "6. MDN Web Docs — Mozilla Developer Network JavaScript, HTML & CSS reference (https://developer.mozilla.org).",
        "7. IETF RFC 7519 — JSON Web Token (JWT) open standard specification."
    ]
    for ref in references:
        p_ref = doc.add_paragraph()
        p_ref.paragraph_format.space_before = Pt(2)
        p_ref.paragraph_format.space_after = Pt(2)
        r_ref = p_ref.add_run(ref)
        r_ref.font.size = Pt(9.5)
        r_ref.font.color.rgb = DARK_SLATE

    # Save
    out_docx_1 = "CareerGraph_Project_Report.docx"
    out_docx_2 = "CareerGraph_Project_Documentation.docx"
    out_docx_3 = "docs/CareerGraph_Technical_Documentation.docx"

    doc.save(out_docx_1)
    doc.save(out_docx_2)
    doc.save(out_docx_3)
    print(f"Successfully generated academic reports: {out_docx_1}, {out_docx_2}, and {out_docx_3}")

if __name__ == "__main__":
    build_academic_report()
