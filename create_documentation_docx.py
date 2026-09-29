import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def create_document():
    doc = Document()

    # -------------------------------------------------------------
    # Page Setup (1 inch margins all around)
    # -------------------------------------------------------------
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)

    # -------------------------------------------------------------
    # Styles & Colors Setup
    # -------------------------------------------------------------
    # Palette
    # Primary Header: #1E3A8A (Navy)
    # Secondary Header: #1E293B (Deep Slate)
    # Accent: #0284C7 (Light Blue / Steel)
    # Body text: #334155 (Slate Charcoal)
    # Table header fill: #1E3A8A
    # Light gray background: #F8FAFC
    # Border color: #CBD5E1

    NAVY = RGBColor(30, 58, 138)
    SLATE = RGBColor(30, 41, 59)
    BODY_COLOR = RGBColor(51, 65, 85)
    MUTED_COLOR = RGBColor(100, 116, 139)

    # Normal Style
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(10.5)
    normal_style.font.color.rgb = BODY_COLOR
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(4)

    # Helper: Set XML background color on cell
    def set_cell_background(cell, hex_color):
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
        cell._tc.get_or_add_tcPr().append(shd)

    # Helper: Set XML margins/padding on cell
    def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = parse_xml(
            f'<w:tcMar {nsdecls("w")}>'
            f'<w:top w:w="{top}" w:type="dxa"/>'
            f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
            f'<w:left w:w="{left}" w:type="dxa"/>'
            f'<w:right w:w="{right}" w:type="dxa"/>'
            f'</w:tcMar>'
        )
        tcPr.append(tcMar)

    # Helper: Set clean borders on table
    def format_table(table, col_widths, border_color="CBD5E1"):
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        tblPr = table._tbl.tblPr
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="6" w:space="0" w:color="{border_color}"/>'
            f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="{border_color}"/>'
            f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="{border_color}"/>'
            f'<w:insideV w:val="none"/>'
            f'<w:left w:val="none"/>'
            f'<w:right w:val="none"/>'
            f'</w:tblBorders>'
        )
        tblPr.append(borders)

        for row_idx, row in enumerate(table.rows):
            # CantSplit to prevent row breaking across pages awkwardly
            trPr = row._tr.get_or_add_trPr()
            trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

            # Repeat header on new page
            if row_idx == 0:
                trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

            for col_idx, cell in enumerate(row.cells):
                cell.width = col_widths[col_idx]
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                set_cell_margins(cell, top=110, bottom=110, left=140, right=140)

                if row_idx == 0:
                    set_cell_background(cell, "1E3A8A")
                    for p in cell.paragraphs:
                        p.paragraph_format.space_before = Pt(2)
                        p.paragraph_format.space_after = Pt(2)
                        for r in p.runs:
                            r.font.bold = True
                            r.font.color.rgb = RGBColor(255, 255, 255)
                            r.font.size = Pt(9.5)
                else:
                    if row_idx % 2 == 1:
                        set_cell_background(cell, "FFFFFF")
                    else:
                        set_cell_background(cell, "F8FAFC")
                    for p in cell.paragraphs:
                        p.paragraph_format.space_before = Pt(1)
                        p.paragraph_format.space_after = Pt(1)
                        for r in p.runs:
                            r.font.size = Pt(9)
                            r.font.color.rgb = BODY_COLOR

    # Helper: Add Callout Box
    def add_callout(text, label="NOTE"):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.cell(0, 0)
        cell.width = Inches(6.5)
        set_cell_background(cell, "F1F5F9")
        set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
        
        # Left accent border
        tcPr = cell._tc.get_or_add_tcPr()
        tcBorders = parse_xml(
            f'<w:tcBorders {nsdecls("w")}>'
            f'<w:left w:val="single" w:sz="24" w:space="0" w:color="0284C7"/>'
            f'<w:top w:val="none"/>'
            f'<w:right w:val="none"/>'
            f'<w:bottom w:val="none"/>'
            f'</w:tcBorders>'
        )
        tcPr.append(tcBorders)
        
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        run_lbl = p.add_run(f"[{label}] ")
        run_lbl.bold = True
        run_lbl.font.color.rgb = RGBColor(2, 132, 199)
        run_lbl.font.size = Pt(9.5)
        run_text = p.add_run(text)
        run_text.font.size = Pt(9.5)
        run_text.font.color.rgb = BODY_COLOR

        p_after = doc.add_paragraph()
        p_after.paragraph_format.space_after = Pt(4)

    # Heading helpers
    def add_h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(16)
        run.font.color.rgb = NAVY
        return p

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(12.5)
        run.font.color.rgb = SLATE
        return p

    def add_h3(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(10.5)
        run.font.color.rgb = RGBColor(71, 85, 105)
        return p

    def add_bullet(bold_prefix, text):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_bold = p.add_run(bold_prefix)
            r_bold.bold = True
            r_bold.font.color.rgb = SLATE
        r_text = p.add_run(text)
        r_text.font.color.rgb = BODY_COLOR
        return p

    # -------------------------------------------------------------
    # Header & Footer (Document-wide)
    # -------------------------------------------------------------
    header = doc.sections[0].header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hrun = hp.add_run("CareerGraph — Technical System Documentation")
    hrun.font.size = Pt(8.5)
    hrun.font.color.rgb = MUTED_COLOR

    footer = doc.sections[0].footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    frun1 = fp.add_run("Page ")
    frun1.font.size = Pt(8.5)
    frun1.font.color.rgb = MUTED_COLOR
    
    # XML Page field
    fldSimple = parse_xml(r'<w:fldSimple %s w:instr="PAGE"/>' % nsdecls('w'))
    fp._p.append(fldSimple)

    # -------------------------------------------------------------
    # COVER / TITLE BLOCK (Clean, professional, not cluttered)
    # -------------------------------------------------------------
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(36)
    title_p.paragraph_format.space_after = Pt(6)
    trun = title_p.add_run("CareerGraph")
    trun.bold = True
    trun.font.size = Pt(28)
    trun.font.color.rgb = NAVY

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_before = Pt(0)
    sub_p.paragraph_format.space_after = Pt(14)
    srun = sub_p.add_run("Personalized IT Career Roadmap Platform with Topological Prerequisite Enforcement and Skill Validation")
    srun.font.size = Pt(13)
    srun.font.color.rgb = SLATE

    # Subtle horizontal line
    div_p = doc.add_paragraph()
    div_p.paragraph_format.space_before = Pt(0)
    div_p.paragraph_format.space_after = Pt(16)
    pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="12" w:space="1" w:color="1E3A8A"/></w:pBdr>')
    div_p._p.get_or_add_pPr().append(pBdr)

    # Meta Table (Clean 2-column key-value)
    meta_table = doc.add_table(rows=6, cols=2)
    meta_widths = [Inches(2.2), Inches(4.3)]
    meta_data = [
        ("Document Type", "System Architecture & Engineering Specification"),
        ("Version & Status", "Version 1.0.0 (Production Verified)"),
        ("Primary Stack", "Vue 3 (Composition API), Node.js, Express, MongoDB Atlas"),
        ("Core Engineering Domains", "Topological Graph Locking, JWT Security, Proof-of-Skill Validation"),
        ("Target Deployments", "Frontend: Vercel SPA (Edge CDN) | Backend: Render (Node.js API)"),
        ("Engineering Team", "Archit Surve, Sanika Vichare")
    ]
    for idx, (label, val) in enumerate(meta_data):
        row = meta_table.rows[idx]
        cell_lbl, cell_val = row.cells[0], row.cells[1]
        cell_lbl.width, cell_val.width = meta_widths[0], meta_widths[1]
        set_cell_background(cell_lbl, "F8FAFC")
        set_cell_background(cell_val, "FFFFFF")
        set_cell_margins(cell_lbl, 80, 80, 100, 100)
        set_cell_margins(cell_val, 80, 80, 100, 100)
        
        pl = cell_lbl.paragraphs[0]
        pl.paragraph_format.space_before = Pt(1)
        pl.paragraph_format.space_after = Pt(1)
        rl = pl.add_run(label)
        rl.bold = True
        rl.font.size = Pt(9.5)
        rl.font.color.rgb = SLATE

        pv = cell_val.paragraphs[0]
        pv.paragraph_format.space_before = Pt(1)
        pv.paragraph_format.space_after = Pt(1)
        rv = pv.add_run(val)
        rv.font.size = Pt(9.5)
        rv.font.color.rgb = BODY_COLOR

    tblPr = meta_table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
        f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
        f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
        f'<w:insideV w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
        f'<w:left w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
        f'<w:right w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

    doc.add_page_break()

    # -------------------------------------------------------------
    # 1. EXECUTIVE SUMMARY & PROBLEM STATEMENT
    # -------------------------------------------------------------
    add_h1("1. Project Overview & Problem Statement")

    doc.add_paragraph(
        "CareerGraph is an interactive, full-stack career roadmap and skill gap platform designed for students, "
        "self-taught programmers, and transitioning professionals in the software engineering sector. "
        "Modern technical education presents an abundance of unstructured learning resources, but lacks clear sequencing, "
        "prerequisite enforcement, and mechanisms to verify practical competency."
    )

    add_h2("1.1 Core Industry Problems Addressed")
    add_bullet("Unstructured Learning Paths: ", "Learners frequently encounter tutorials and courses without understanding necessary prerequisites. For instance, attempting Docker and Kubernetes before mastering Linux shell fundamentals or networking concepts leads to high dropout rates.")
    add_bullet("Static, Non-Interactive Visualizations: ", "Traditional roadmap references (such as static images or read-only charts) do not account for a learner's prior experience or real-time progress.")
    add_bullet("Fragmented Skill Portability: ", "A learner who masters Git or Python in a web development roadmap must often restart from scratch when switching to an adjacent track such as DevOps or Data Engineering.")
    add_bullet("Unverified Competency Claims: ", "Self-assessment alone produces inaccurate skill metrics. Platforms typically offer either complete honor-system checkmarks or rigid, high-friction examination portals.")

    add_h2("1.2 The CareerGraph Solution")
    doc.add_paragraph(
        "CareerGraph organizes curricula into a cohesive five-tier hierarchy: "
        "Career Track → Skills → Courses → Modules → Lessons. "
        "The system evaluates the learner's incoming knowledge during onboarding, dynamically calculates their remaining "
        "skill gap, locks advanced phases until prerequisite competencies are completed, and verifies knowledge through a "
        "server-graded Proof of Skill validation engine."
    )

    add_callout(
        "Universal Skill Portability: Because skills are identified by normalized universal keys (e.g., 'git', 'python', 'sql'), "
        "progress attained in one career track immediately propagates to any adjacent track requiring the same competency.",
        label="CORE PRINCIPLE"
    )

    # -------------------------------------------------------------
    # 2. SYSTEM ARCHITECTURE
    # -------------------------------------------------------------
    add_h1("2. System Architecture & High-Level Design")

    doc.add_paragraph(
        "CareerGraph follows a decoupled client-server architecture composed of three primary layers: "
        "the Presentation Tier (Vue 3 Single Page Application), the Application Tier (Node.js & Express REST API), "
        "and the Persistence Tier (MongoDB with Mongoose ODM)."
    )

    # Embed Figure 1
    if os.path.exists("docs/system_architecture.png"):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(8)
        p_img.paragraph_format.space_after = Pt(2)
        doc.add_picture("docs/system_architecture.png", width=Inches(6.2))
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_after = Pt(12)
        rcap = p_cap.add_run("Figure 1: High-Level System Architecture and Tier Decomposition")
        rcap.font.size = Pt(8.5)
        rcap.font.italic = True
        rcap.font.color.rgb = MUTED_COLOR

    add_h2("2.1 Layer Breakdown")
    add_bullet("Presentation Tier (Frontend): ", "Built using Vue 3 with the Composition API (<script setup>) and Vite. Centralized state is handled via Pinia stores (auth, roadmap, user). The interactive roadmap view features an SVG canvas calculating cubic Bézier connector curves, pan/zoom viewports, single-click status toggles, and contextual topic drawers.")
    add_bullet("Application Tier (Backend API): ", "Constructed on Node.js with Express.js. Implements stateless JWT authentication, password hashing with bcrypt, algorithmic dependency resolution (skillGapService.js), and server-side quiz grading (quizController.js).")
    add_bullet("Persistence Tier (Database): ", "Hosted on MongoDB Atlas (or local replica set), managed through Mongoose 8.5. Nine normalized models with compound uniqueness indexes guarantee sub-millisecond query performance and transactional idempotency.")

    add_h2("2.2 Technology Stack Justification")
    tech_table = doc.add_table(rows=7, cols=3)
    tech_widths = [Inches(1.8), Inches(1.8), Inches(2.9)]
    tech_data = [
        ("Layer / Component", "Technology Selected", "Design Rationale & Engineering Benefit"),
        ("Frontend Framework", "Vue 3 (Composition API)", "Clean reactivity model, fine-grained component encapsulation, and high rendering performance for SVG nodes."),
        ("State Management", "Pinia 2.1", "Type-safe modular stores without boilerplate; direct reactivity for user session, active roadmap, and profile metrics."),
        ("Backend Framework", "Node.js & Express 4.19", "Lightweight, event-driven I/O model well suited for concurrent API requests and async MongoDB queries."),
        ("Database & ODM", "MongoDB + Mongoose 8.5", "Flexible JSON document storage matches hierarchical curriculum data (Levels, Modules, Lessons) naturally."),
        ("Authentication", "Stateless JWT + bcrypt", "Eliminates server-side session overhead; permits horizontal scaling and seamless client rehydration."),
        ("Frontend Build Tool", "Vite 5.3", "Instant development server startup via ES modules, optimized Rollup production bundling, and fast HMR.")
    ]
    for r_idx, row in enumerate(tech_table.rows):
        for c_idx, cell in enumerate(row.cells):
            cell.text = tech_data[r_idx][c_idx]
    format_table(tech_table, tech_widths)

    doc.add_page_break()

    # -------------------------------------------------------------
    # 3. DATABASE SCHEMAS & DATA MODELING
    # -------------------------------------------------------------
    add_h1("3. Database Schema & Data Dictionary")

    doc.add_paragraph(
        "The persistence layer utilizes nine distinct Mongoose schemas. "
        "The model balances normalization (avoiding duplicated curriculum content across user progress records) "
        "with controlled embedding (nesting modules and lessons inside level documents to reduce read join overhead)."
    )

    # Embed Figure 2
    if os.path.exists("docs/er_diagram.png"):
        p_img2 = doc.add_paragraph()
        p_img2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img2.paragraph_format.space_before = Pt(8)
        p_img2.paragraph_format.space_after = Pt(2)
        doc.add_picture("docs/er_diagram.png", width=Inches(6.2))
        
        p_cap2 = doc.add_paragraph()
        p_cap2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap2.paragraph_format.space_after = Pt(12)
        rcap2 = p_cap2.add_run("Figure 2: Entity-Relationship Diagram and Referential Cardinality")
        rcap2.font.size = Pt(8.5)
        rcap2.font.italic = True
        rcap2.font.color.rgb = MUTED_COLOR

    add_h2("3.1 Core Entity Dictionary")

    # Table of Entities
    entities = [
        ("User", "Core user account credentials, hashed password, role authorization (user/admin), target career reference, and onboarding completion flag."),
        ("UserProfile", "1:1 user profile containing academic stream, degree, graduation year, employment status, and declared skill proficiency ratings (0 to 4)."),
        ("Career", "IT career paths (Full Stack, DevOps, AI/ML, Frontend, Backend, Data Analyst) with required skills mapped to phases (Foundations, Core, Advanced)."),
        ("Skill", "Atomic skill definitions with universal identifiers (skillId), descriptions, industry use cases, external resources, and self-referencing prerequisites."),
        ("Course", "Educational curriculum associated 1:1 with a Skill, providing the structured syllabus divided into difficulty levels."),
        ("Level", "Curriculum tiers (Beginner, Intermediate, Advanced) containing embedded modules and granular lessons with code demonstrations."),
        ("CourseProgress", "Aggregated completion percentage (0-100%) for a user on a given course/skill, tracking completed vs total lesson count and verification badge status."),
        ("LessonProgress", "Atomic completion log recording exactly which user completed which lesson, indexed by compound uniqueness to prevent duplicates."),
        ("SkillQuiz", "Validation quiz associated with a skill containing at least 3 multiple-choice questions, options, correct answers, and instructional explanations.")
    ]
    entity_table = doc.add_table(rows=len(entities) + 1, cols=2)
    entity_widths = [Inches(1.8), Inches(4.7)]
    entity_table.rows[0].cells[0].text = "Entity Model"
    entity_table.rows[0].cells[1].text = "Description & Architectural Role"
    for idx, (ent, desc) in enumerate(entities):
        row = entity_table.rows[idx + 1]
        row.cells[0].text = ent
        row.cells[1].text = desc
    format_table(entity_table, entity_widths)

    add_h2("3.2 Detailed Field Specifications")
    
    # User & UserProfile
    add_h3("A. User and UserProfile Models")
    user_fields = [
        ("Field Name", "Type", "Constraints", "Description"),
        ("_id", "ObjectId", "Primary Key", "System generated unique identifier."),
        ("name", "String", "Required, Trim", "User full display name."),
        ("email", "String", "Required, Unique, Lowercase", "Unique login credential."),
        ("password", "String", "Required, select: false", "Bcrypt hashed password (10 salt rounds)."),
        ("targetCareer", "ObjectId", "Ref: Career, Optional", "Active career roadmap being tracked."),
        ("onboardingComplete", "Boolean", "Default: false", "Flag indicating baseline assessment completion."),
        ("profile.education", "Subdocument", "Optional", "Academic background: degree, field, graduation year."),
        ("profile.skills", "Array", "Embedded subdocuments", "List of declared skills: { skill: ObjectId, level: 0..4, verified: Boolean }.")
    ]
    t_user = doc.add_table(rows=len(user_fields), cols=4)
    format_table(t_user, [Inches(1.5), Inches(1.2), Inches(1.5), Inches(2.3)])
    for r_idx, row in enumerate(t_user.rows):
        for c_idx, cell in enumerate(row.cells):
            cell.text = user_fields[r_idx][c_idx]

    # Skill, Career & Course
    add_h3("B. Career, Skill and Course Models")
    curric_fields = [
        ("Field Name", "Type", "Constraints", "Description"),
        ("Career.requiredSkills", "Array", "Embedded", "Required skills with requiredLevel and assigned phase (foundations, core, advanced)."),
        ("Skill.skillId", "String", "Required, Unique, Index", "Universal cross-track slug (e.g., 'git', 'docker', 'python')."),
        ("Skill.prerequisites", "Array<ObjectId>", "Ref: Skill", "Self-referencing prerequisite skills required before this skill."),
        ("Course.skill", "ObjectId", "Ref: Skill, Unique (1:1)", "Direct reference linking curriculum to the corresponding Skill."),
        ("Level.modules", "Array", "Embedded", "Modules containing lessons with title, description, content, and codeExample.")
    ]
    t_curric = doc.add_table(rows=len(curric_fields), cols=4)
    format_table(t_curric, [Inches(1.7), Inches(1.2), Inches(1.4), Inches(2.2)])
    for r_idx, row in enumerate(t_curric.rows):
        for c_idx, cell in enumerate(row.cells):
            cell.text = curric_fields[r_idx][c_idx]

    # Progress & Quizzes
    add_h3("C. Progress and Validation Models")
    prog_fields = [
        ("Field Name", "Type", "Constraints", "Description"),
        ("CourseProgress.user", "ObjectId", "Ref: User, Compound UK", "Owner of progress record; unique across (user, course)."),
        ("CourseProgress.percent", "Number", "0 to 100, Default: 0", "Overall course completion percentage."),
        ("CourseProgress.verified", "Boolean", "Default: false", "Flag set to true when user passes the Proof of Skill quiz."),
        ("LessonProgress", "Collection", "Compound UK: (user, lesson)", "Atomic completion record. Guarantees idempotent toggles."),
        ("SkillQuiz.questions", "Array", "Min 3 questions", "Validation questions with options, correctIndex, and explanation.")
    ]
    t_prog = doc.add_table(rows=len(prog_fields), cols=4)
    format_table(t_prog, [Inches(1.7), Inches(1.2), Inches(1.4), Inches(2.2)])
    for r_idx, row in enumerate(t_prog.rows):
        for c_idx, cell in enumerate(row.cells):
            cell.text = prog_fields[r_idx][c_idx]

    doc.add_page_break()

    # -------------------------------------------------------------
    # 4. CORE ALGORITHMS & DOMAIN LOGIC
    # -------------------------------------------------------------
    add_h1("4. Core Algorithms & Domain Business Logic")

    doc.add_paragraph(
        "CareerGraph implements algorithmic business rules to evaluate user knowledge, "
        "enforce curriculum order, and ensure that skill verification is meaningful."
    )

    add_h2("4.1 Dynamic Roadmap Assembly (skillGapService.js)")
    doc.add_paragraph(
        "When an authenticated user requests their roadmap, the backend does not return a static template. "
        "Instead, it dynamically generates the graph:"
    )
    add_bullet("Curriculum Lookup: ", "Loads the target Career and retrieves all required skills, populated with prerequisite references.")
    add_bullet("Baseline Merging: ", "Retrieves the user's declared knowledge ratings from UserProfile.skills.")
    add_bullet("Progress Correlation: ", "Fetches actual lesson completions from CourseProgress by matching course IDs, skill ObjectIds, or universal skillId strings.")
    add_bullet("Status Resolution: ", "Assigns node status based on dual criteria: a node is 'completed' (100%) if course progress is 100% OR user declared level >= 3; 'in_progress' if progress > 0% or declared level > 0; and 'not_started' otherwise.")

    add_h2("4.2 Topological Prerequisite & Phase Locking Engine")
    doc.add_paragraph(
        "To prevent skipping fundamental knowledge, computeUnlockedPhases executes a two-level locking routine:"
    )
    add_bullet("Phase-Level Gating: ", "Nodes are divided into Phase 1 (Foundations), Phase 2 (Core Stack), and Phase 3 (Advanced & Ecosystem). Phase 1 is always unlocked. Phase 2 unlocks only after every skill in Phase 1 is marked 'completed'. Phase 3 unlocks only after every skill in Phase 2 is marked 'completed'.")
    add_bullet("Node-Level Dependency Checking: ", "Even inside an unlocked phase, individual nodes verify their specific prerequisites. If Skill B requires Skill A, Skill B remains locked until Skill A is completed.")
    add_bullet("Next Recommended Skill Pointer: ", "The engine traverses unlocked nodes in topological order and assigns the 'recommended' badge to the first non-completed skill, giving the learner immediate direction.")

    add_h2("4.3 Proof-of-Skill Validation Engine (quizController.js)")
    doc.add_paragraph(
        "To establish authentic competency without relying solely on manual checkmarks, CareerGraph incorporates a "
        "server-side quiz validation engine:"
    )
    add_bullet("Advanced Phase Gating: ", "The quiz cannot be unlocked prematurely. A user must reach the Advanced phase (completing all Beginner and Intermediate lessons, or having an existing proficiency level >= 3) before access is granted.")
    add_bullet("Payload Sanitization: ", "When fetching questions via GET /api/skills/:skillId/quiz, the server explicitly strips 'correctIndex' and 'explanation' fields from the payload. This prevents client-side answer key inspection via browser DevTools.")
    add_bullet("Server-Side Grading: ", "Answers are submitted via POST /api/skills/:skillId/quiz/submit. The server compares user answers against stored correct indices. A score of at least 2 out of 3 (>= 66.7%) is required to pass.")
    add_bullet("Verification Rollup: ", "On passing, the system automatically: (1) marks all lessons in the course as complete, (2) updates CourseProgress to 100% with verified = true and the quiz score, and (3) elevates the user's skill profile rating to Advanced with a Verified status.")

    add_callout(
        "Security Integrity: Answer evaluation occurs strictly in the Node.js application layer. "
        "Even if a user manipulates network responses, the backend database write only occurs upon passing the server check.",
        label="SECURITY GUARANTEE"
    )

    doc.add_page_break()

    # -------------------------------------------------------------
    # 5. FRONTEND ARCHITECTURE & INTERACTIVE CANVAS
    # -------------------------------------------------------------
    add_h1("5. Frontend Architecture & Interactive Canvas")

    doc.add_paragraph(
        "The client application is built with Vue 3 and Vite, focusing on immediate responsiveness and smooth graph navigation."
    )

    add_h2("5.1 Interactive SVG Roadmap Canvas (Roadmap.vue)")
    doc.add_paragraph(
        "The primary learning view offers both an interactive Canvas Graph and a structured Phase Flow:"
    )
    add_bullet("Dynamic Bézier Connector Curves: ", "The canvas calculates cubic Bézier paths between prerequisite and dependent nodes in real time. Marker arrows visually illustrate dependency flow across phase columns.")
    add_bullet("Pan and Zoom Viewport: ", "Supports drag-to-pan, mouse wheel zooming (with clamp limits from 50% to 150%), and a floating HUD with Zoom In, Zoom Out, and Center & Reset controls.")
    add_bullet("Quick Status Toggle Badge: ", "Clicking directly on a node's status badge cycles through Not Started ➔ In Progress ➔ Completed in one click, updating the backend optimistically.")
    add_bullet("Right-Click Context Menu: ", "Right-clicking any node opens a floating action menu to jump directly to any status or navigate to detailed lessons.")
    add_bullet("Prerequisite Lock Toast: ", "Attempting to toggle a locked skill displays an informative banner detailing unmet prerequisites.")

    add_h2("5.2 Slide-Out Topic Preview Drawer (SkillPreviewDrawer.vue)")
    doc.add_paragraph(
        "Clicking any node slides open an in-canvas preview drawer without navigating away from the roadmap. "
        "The drawer includes:"
    )
    add_bullet("Skill Overview: ", "Conceptual definition, industry relevance ('Why It Matters'), and practical applications.")
    add_bullet("Curated Learning Materials: ", "Direct links to official documentation, interactive sandboxes, and video tutorials.")
    add_bullet("Interactive Curriculum Checklist: ", "Beginner, Intermediate, and Advanced module lists with individual lesson checkboxes that update course progress in real time.")
    add_bullet("Proof of Skill Launcher: ", "Launches the SkillQuizModal when the user reaches the Advanced phase.")

    add_h2("5.3 In-Place Career Switcher (CareerSwitchModal.vue)")
    doc.add_paragraph(
        "Allows users to explore and switch career paths directly from the roadmap. "
        "Features real-time filtering by academic stream (Computer Science, IT, AI & Data Science) "
        "and degree type (B.Tech, BCA, B.Sc). Progress across overlapping skills is preserved seamlessly."
    )

    add_h2("5.4 State Management with Pinia")
    doc.add_paragraph(
        "Application state is separated into three specialized Pinia stores:"
    )
    add_bullet("auth.js: ", "Manages user authentication tokens, login/register actions, and session persistence in localStorage.")
    add_bullet("roadmap.js: ", "Stores active career graph nodes, phase allocations, overall completion percentage, and active drawer state.")
    add_bullet("user.js: ", "Maintains user profile attributes, declared skill levels, and academic degree preferences.")

    doc.add_page_break()

    # -------------------------------------------------------------
    # 6. RESTFUL API SPECIFICATION
    # -------------------------------------------------------------
    add_h1("6. RESTful API Specifications")

    doc.add_paragraph(
        "All API endpoints follow REST conventions, communicating via JSON payloads. "
        "Endpoints requiring authentication expect an HTTP header: Authorization: Bearer <token>."
    )

    api_endpoints = [
        ("Method", "Endpoint URI", "Auth", "Payload / Query", "Response Status & Description"),
        ("POST", "/api/auth/register", "Public", "{ name, email, password }", "201 Created — Returns user profile and signed JWT."),
        ("POST", "/api/auth/login", "Public", "{ email, password }", "200 OK — Returns authenticated user and signed JWT."),
        ("GET", "/api/auth/me", "Bearer JWT", "None", "200 OK — Current authenticated user profile."),
        ("GET", "/api/careers", "Optional", "?stream=&degree=&category=", "200 OK — Filtered list of career roadmap tracks."),
        ("GET", "/api/careers/:slug", "Optional", "None", "200 OK — Career details with populated required skills."),
        ("GET", "/api/skills/:skillId", "Optional", "None", "200 OK — Skill details, prerequisites, and resource links."),
        ("GET", "/api/skills/:skillId/quiz", "Bearer JWT", "None", "200 OK — 3 sanitized validation questions (answers omitted)."),
        ("POST", "/api/skills/:skillId/quiz/submit", "Bearer JWT", "{ answers: [0, 2, 1] }", "200 OK — Grading results, pass status, and verified badge."),
        ("GET", "/api/roadmap", "Bearer JWT", "None", "200 OK — Dynamic roadmap for active career with phase locks."),
        ("GET", "/api/roadmap/:careerId", "Bearer JWT", "None", "200 OK — Dynamic roadmap for specified career track."),
        ("PUT", "/api/progress/skill/:id/status", "Bearer JWT", "{ status: 'completed' }", "200 OK — Updated skill node status and rollup metrics."),
        ("POST", "/api/progress/toggle-lesson", "Bearer JWT", "{ lessonId, courseId }", "200 OK — Updated lesson state and course percentage."),
        ("PUT", "/api/users/target-career", "Bearer JWT", "{ careerId }", "200 OK — Updated active career path."),
        ("GET", "/api/users/profile", "Bearer JWT", "None", "200 OK — Detailed profile, degree, and verified badges.")
    ]

    t_api = doc.add_table(rows=len(api_endpoints), cols=5)
    api_widths = [Inches(0.8), Inches(1.8), Inches(0.9), Inches(1.4), Inches(1.6)]
    for r_idx, row in enumerate(t_api.rows):
        for c_idx, cell in enumerate(row.cells):
            cell.text = api_endpoints[r_idx][c_idx]
    format_table(t_api, api_widths)

    add_callout(
        "Standardized Error Handling: Unhandled exceptions and validation errors return a uniform JSON format: "
        "{ message: 'Descriptive error string', statusCode: 4xx/5xx }.",
        label="API CONVENTION"
    )

    doc.add_page_break()

    # -------------------------------------------------------------
    # 7. SECURITY, AUTHENTICATION & DATA PROTECTION
    # -------------------------------------------------------------
    add_h1("7. Security, Authentication & Data Protection")

    doc.add_paragraph(
        "Security practices are applied across the presentation, application, and persistence tiers:"
    )

    add_bullet("Password Encryption: ", "Passwords are never stored in plaintext. Mongoose pre-save hooks hash passwords using bcryptjs with 10 salt rounds before persistence.")
    add_bullet("Stateless JWT Authorization: ", "API routes are protected via authMiddleware.js. Each request validates the signature and expiration of the Bearer token without requiring database session lookups.")
    add_bullet("Projection Defense: ", "The password hash is explicitly configured with select: false in the User schema. Controller responses sanitize user objects via toSafeObject() to prevent accidental credential leakage.")
    add_bullet("Quiz Inspection Prevention: ", "Correct answer indices and instructional explanations are removed from quiz retrieval payloads, preventing client-side examination cheating.")
    add_bullet("Database Concurrency & Idempotency: ", "Compound unique indexes ({ user: 1, course: 1 } and { user: 1, lesson: 1 }) ensure that rapid concurrent clicks or network retries do not generate duplicate progress records.")

    # -------------------------------------------------------------
    # 8. MODULE BREAKDOWN & WORKLOAD DISTRIBUTION
    # -------------------------------------------------------------
    add_h1("8. Engineering Workload Distribution")

    doc.add_paragraph(
        "For project defense and academic evaluation, the engineering responsibilities are distributed across "
        "six specialized functional modules:"
    )

    workload = [
        ("Module / Role", "Primary Technical Ownership", "Specific Deliverables & Code Artifacts"),
        ("Module 1: Authentication & Security", "User Identity & Access Architecture", "• User and UserProfile Mongoose schemas\n• Bcrypt hashing & pre-save sanitization hooks\n• JWT route protection middleware (authMiddleware.js)\n• Profile management & authentication controllers"),
        ("Module 2: Graph Algorithms", "Topological Prerequisite Engine", "• Graph resolution service (skillGapService.js)\n• Three-tier phase organization & prerequisite validation\n• Dynamic locking logic & next-skill recommendation\n• Universal cross-track skill calculation"),
        ("Module 3: Curriculum & Database", "Database Modeling & Content Seeding", "• Normalized schemas (Career, Skill, Course, Level, Quiz)\n• Comprehensive database seeding script (seed.js)\n• MongoDB compound indexing for sub-millisecond lookups\n• Academic stream and degree category mapping"),
        ("Module 4: Canvas Visualization", "Interactive Roadmap Graph", "• Interactive canvas component (Roadmap.vue)\n• SVG cubic Bézier curve calculation engine\n• Quick status toggle badge system & right-click menu\n• Viewport pan, drag, and zoom HUD controls"),
        ("Module 5: UI/UX & Drawer Components", "Workflow Modals & Topic Inspection", "• Multi-step assessment wizard (Onboarding.vue)\n• Slide-out topic preview drawer (SkillPreviewDrawer.vue)\n• In-place career switcher (CareerSwitchModal.vue)\n• Proof-of-Skill validation dialog (SkillQuizModal.vue)"),
        ("Module 6: State & Deployment", "Pinia Integration & DevOps", "• Centralized stores (auth.js, roadmap.js, user.js)\n• Axios API service with automatic JWT interceptors\n• Vercel SPA routing configuration (vercel.json)\n• Render backend web service & production verification")
    ]

    t_work = doc.add_table(rows=len(workload), cols=3)
    work_widths = [Inches(1.8), Inches(1.8), Inches(2.9)]
    for r_idx, row in enumerate(t_work.rows):
        for c_idx, cell in enumerate(row.cells):
            cell.text = workload[r_idx][c_idx]
    format_table(t_work, work_widths)

    doc.add_page_break()

    # -------------------------------------------------------------
    # 9. SETUP, EXECUTION & DEPLOYMENT GUIDE
    # -------------------------------------------------------------
    add_h1("9. Setup, Execution & Operational Guide")

    doc.add_paragraph(
        "Follow these steps to configure and run CareerGraph locally or deploy to production environments."
    )

    add_h2("9.1 Backend Setup")
    doc.add_paragraph("Navigate to the backend directory and install dependencies:")
    add_bullet("Command: ", "cd backend && npm install")
    doc.add_paragraph("Create a backend/.env configuration file with the following variables:")
    add_bullet("PORT: ", "5001")
    add_bullet("MONGO_URI: ", "mongodb://localhost:27017/careergraph (or MongoDB Atlas URI)")
    add_bullet("JWT_SECRET: ", "your_super_secret_jwt_key")
    doc.add_paragraph("Seed the database with initial career tracks, skills, and quizzes, then launch the API:")
    add_bullet("Seed Command: ", "npm run seed")
    add_bullet("Development Server: ", "npm run dev  (Server listens at http://localhost:5001)")

    add_h2("9.2 Frontend Setup")
    doc.add_paragraph("Navigate to the frontend directory and install dependencies:")
    add_bullet("Command: ", "cd frontend && npm install")
    doc.add_paragraph("Create a frontend/.env configuration file:")
    add_bullet("VITE_API_BASE_URL: ", "http://localhost:5001/api")
    doc.add_paragraph("Launch the Vite development server:")
    add_bullet("Development Server: ", "npm run dev  (Application accessible at http://localhost:5173)")
    add_bullet("Production Build: ", "npm run build  (Outputs optimized static assets to frontend/dist)")

    add_h2("9.3 Production Deployment Architecture")
    add_bullet("Frontend on Vercel: ", "Deployed as a Single Page Application. Routing is managed by frontend/vercel.json with a catch-all rewrite rule (/(.*) -> /) ensuring client-side routes (like /roadmap and /onboarding) resolve properly without 404 errors.")
    add_bullet("Backend on Render: ", "Deployed as a Node.js Web Service running 'npm start' with environment variables set in the Render dashboard. Auto-deploy triggers on every push to the GitHub main branch.")
    add_bullet("Database on MongoDB Atlas: ", "Managed cloud cluster with automated daily backups, IP access controls, and connection pooling.")

    # -------------------------------------------------------------
    # Save Documents
    # -------------------------------------------------------------
    out_path_docs = "docs/CareerGraph_Technical_Documentation.docx"
    out_path_root = "CareerGraph_Project_Documentation.docx"
    
    doc.save(out_path_docs)
    doc.save(out_path_root)
    print(f"Successfully generated {out_path_docs} and {out_path_root}")

if __name__ == "__main__":
    create_document()
